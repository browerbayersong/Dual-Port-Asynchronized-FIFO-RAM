"""Audit measured report provenance, then capture the delivered file set after visual QA."""
import argparse
import ast
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]


def digest(path,binary=False):
    data=path.read_bytes()
    return hashlib.sha256(data if binary else data.replace(b'\r\n',b'\n')).hexdigest()


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--visual-qa-pages',type=int,required=True,help='Actual pages visually inspected by the operator')
    args=ap.parse_args()
    old=json.loads((ROOT/'reports/analysis_manifest.json').read_text())
    for rel,expected in old['source_sha256_lf'].items():
        if rel.split('/')[0] in ['rtl','tb','sim']:
            assert digest(ROOT/rel)==expected, 'Previously verified source changed: '+rel
    impl=json.loads((ROOT/'reports/implementation/summary.json').read_text())
    inputs=json.loads((ROOT/'reports/implementation/manifest.json').read_text())
    for rel,expected in inputs['source_sha256_lf'].items():
        assert digest(ROOT/rel)==expected, 'Target implementation source changed: '+rel
    configs=impl['configurations']
    assert [c['sync_stages'] for c in configs]==[2,3]
    assert configs[1]['registers']-configs[0]['registers']==44
    for c in configs:
        assert c['lut6']==1040 and c['eram20k']==2 and c['io']==231
        assert c['interface']['setup_wns_ns']<0 and c['interface']['hold_wns_ns']<0
        assert c['core']['setup_wns_ns']>0 and c['core']['hold_wns_ns']>0
        assert len(c['interface']['clocks'])==5 and len(c['core']['clocks'])==5
        assert c['cdc_max_delay_valid_paths']==472 and len(c['cdc_bounds'])==16
        assert all(b['dominated']==b['total'] and b['shadowed']==b['ignored']==0 for b in c['cdc_bounds'])
        assert len(c['gray_delay_evidence'])==4 and all(g['reported_bits']==9 for g in c['gray_delay_evidence'])
        assert c['power_w'] is None
        flow=(ROOT/f'reports/implementation/s{c["sync_stages"]}_flow.log').read_text()
        assert 'TMR-5506' in flow and 'IMPLEMENTATION: PASS' in flow
        assert 'DeriveClock' not in (ROOT/f'reports/implementation/s{c["sync_stages"]}_route.timing').read_text()
    power=json.loads((ROOT/'reports/implementation/power_status.json').read_text())
    assert power['power_w'] is None and power['status']=='unavailable_missing_device_power_library'
    assert 'PWR-5001' in (ROOT/'reports/implementation/power_attempt.log').read_text()
    repro=json.loads((ROOT/'reports/reproduction/summary.json').read_text())
    assert repro['single_configuration']=='PASS' and repro['full_regression']=='PASS 10/10'
    assert repro['human_independent_reproduction'] is False
    for rel,expected in repro['copied_verification_source_sha256_lf'].items():assert digest(ROOT/rel)==expected
    assert (ROOT/'reports/reproduction/full_regression.log').read_text().count('RESULT: PASS')==10
    for stage in [2,3]:
        log=(ROOT/f'reports/implementation/s{stage}_wrapper_verify.log').read_text()
        assert 'RESULT: PASS' in log and 'CHECK_FAIL' not in log
    for p in (ROOT/'tools').glob('*.py'):ast.parse(p.read_text(encoding='utf-8'),filename=p.name)
    changed_docs=[ROOT/'README.md',ROOT/'constraints/README.md']+list((ROOT/'docs').glob('*.md'))
    changed_docs+=[ROOT/'reports/implementation/REPORT.md',ROOT/'reports/mtbf/REPORT.md']
    for p in changed_docs:
        for target in re.findall(r'\]\(([^)]+)\)',p.read_text(encoding='utf-8')):
            if re.match(r'\w+://|mailto:|#',target):continue
            assert (p.parent/target.split('#')[0]).exists(), f'Broken local link: {p.name}: {target}'
    doc=Document(ROOT/'deliverables/双异步FIFO课程设计报告.docx')
    assert len(doc.inline_shapes)==6 and len(doc.tables)==16
    expected_sections=['1 设计题目','2 设计目的','3 设计内容','4 设计结果分析','5 总结','6 参考文献']
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 1']==expected_sections
    anchors=doc._element.xpath('.//wp:anchor')
    assert len(anchors)==1
    cover=json.loads((ROOT/'assets/course-report-cover.json').read_text(encoding='utf-8'))
    assert digest(ROOT/'assets/course-report-cover.pdf',True)==cover['extracted_pdf_sha256']
    assert digest(ROOT/'assets/course-report-cover.png',True)==cover['cover_png_sha256']
    assert len(PdfReader(ROOT/'assets/course-report-cover.pdf').pages)==1
    assert cover['page_content_stream_unchanged'] is True
    with zipfile.ZipFile(ROOT/'deliverables/双异步FIFO课程设计报告.docx') as package:
        media=[package.read(p) for p in package.namelist() if p.startswith('word/media/')]
        assert cover['cover_png_sha256'] in [hashlib.sha256(data).hexdigest() for data in media]
    assert all(not c.text for row in doc.tables[-1].rows[1:] for c in row.cells)
    assert not doc._element.xpath('.//w:pBdr') and not doc.styles.element.xpath('.//w:pBdr')
    qa=Path('D:/fifo_build/report_qa')
    pdf=next(qa.glob('*.pdf'));reader=PdfReader(pdf)
    assert len(reader.pages)==args.visual_qa_pages
    text='\n'.join(p.extract_text() for p in reader.pages)
    assert '\ufffd' not in text
    for token in ['TMR-5506','PWR-5001','0.929','1.203','1532','1576','1040','MTBF']:
        assert token in text, 'Missing expected report evidence: '+token
    sources={}
    for directory in ['rtl','tb','sim','analysis','constraints','tools','docs']:
        for p in sorted((ROOT/directory).rglob('*')):
            if p.is_file() and p.suffix in ['.v','.sv','.f','.ps1','.py','.cpp','.txt','.sdc','.tcl','.md']:
                sources[p.relative_to(ROOT).as_posix()]=digest(p)
    sources['README.md']=digest(ROOT/'README.md')
    artifacts={}
    for directory in ['reports/implementation','reports/figures','reports/reproduction','reports/mtbf','deliverables','assets']:
        for p in sorted((ROOT/directory).rglob('*')):
            if p.is_file():
                binary=p.suffix in ['.png','.pdf','.docx']
                artifacts[p.relative_to(ROOT).as_posix()]={'sha256':digest(p,binary),'encoding':'raw_bytes' if binary else 'canonical_LF_bytes'}
    manifest={'captured_at':datetime.now(timezone(timedelta(hours=8))).isoformat(),
        'checkout_parent_baseline':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
        'commit_note':'The containing Git commit identifies the delivered set; hashes identify exact measured inputs and outputs. Earlier reports retain their original manifests.',
        'scope':'Software target analysis and course report; full IO timing is not closed, TMR-5506 remains, power W unavailable',
        'cover_source':cover,'report_top_level_sections':expected_sections,
        'visual_qa':{'renderer':'WPS COM PDF export plus packaged DOCX rasterizer and Poppler',
            'pages_inspected':args.visual_qa_pages,'qa_pdf_sha256':digest(pdf,True),
            'page_png_sha256':{f'page-{i}.png':digest(qa/f'page-{i}.png',True) for i in range(1,args.visual_qa_pages+1)}},
        'source_sha256_lf':sources,'artifacts':artifacts,
        'checks':{'previous_rtl_tb_sim_unchanged':True,'implementation_inputs_match':True,'clean_reproduction_10_of_10':True,
            'wrapper_checks_both_stages':True,'five_clocks_and_472_bounds':True,'negative_interface_slack_preserved':True,
            'missing_power_preserved':True,'personal_fields_blank':True,'six_report_figures':True,'local_document_links':True,
            'original_pdf_cover_preserved':True,'six_required_top_level_sections':True,'sixteen_technical_and_division_tables':True}}
    (ROOT/'reports/delivery_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print('DELIVERY_AUDIT: PASS',args.visual_qa_pages,'pages, original RTL/TB unchanged, target/reproduction/report consistent')


if __name__=='__main__':main()
